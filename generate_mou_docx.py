#!/usr/bin/env python3
"""Generate a Word (.docx) version of the Telom-X-Gene supply MOU."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Base font
style = doc.styles["Normal"]
style.font.name = "Malgun Gothic"
style.font.size = Pt(10.5)
# Ensure East Asian font applies
from docx.oxml.ns import qn
style.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def heading(text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = "Malgun Gothic"
    run.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    return h


def para(text, bold=False, align=None, size=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


# Title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("양해각서\n(Memorandum of Understanding)")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("Telom-X-Gene 공급에 관한 양해각서")
rs.bold = True
rs.font.size = Pt(13)

doc.add_paragraph()
para("본 양해각서(이하 \"본 MOU\")는 아래의 당사자 간에 Telom-X-Gene(이하 \"본 제품\")의 "
     "공급 및 협력에 관한 기본적 사항을 정하기 위하여 체결된다.")

# Parties table
t = doc.add_table(rows=3, cols=2)
t.style = "Table Grid"
t.alignment = WD_TABLE_ALIGNMENT.CENTER
cells = t.rows[0].cells
cells[0].text = "구분"
cells[1].text = "당사자"
r1 = t.rows[1].cells
r1[0].text = "갑 (공급자)"
r1[1].text = "TelomX 사 (Telom-X-Gene)\n주소: [본사 주소]\n대표자: [성명]"
r2 = t.rows[2].cells
r2[0].text = "을 (구매자)"
r2[1].text = "[중국 제약회사명]\n주소: [중국 주소]\n대표자: [성명]"

para("(이하 갑과 을을 개별적으로 \"당사자\", 총칭하여 \"당사자들\"이라 한다.)")

# Sections
sections = [
    ("전문 (Preamble)", [
        "갑은 텔로미어(telomere) 기반 유전자 치료/원료 기술인 Telom-X-Gene의 연구·개발 및 생산 역량을 "
        "보유하고 있으며, 을은 중화인민공화국(이하 \"중국\") 내에서 의약품의 제조·유통·판매에 관한 사업을 영위하고 있다.",
        "당사자들은 을의 중국 시장을 대상으로 한 Telom-X-Gene의 안정적 공급 및 상호 협력 관계 구축에 "
        "상호 관심을 가지고 있으며, 향후 정식 공급계약(이하 \"본계약\") 체결을 위한 기본 원칙과 협력 의사를 "
        "확인하고자 본 MOU를 체결한다.",
    ]),
    ("제1조 (목적)", [
        "본 MOU는 갑이 을에게 Telom-X-Gene을 공급하고 을이 이를 중국 시장에 공급·활용함에 있어 필요한 "
        "협력의 기본 방향과 원칙을 정하고, 향후 본계약 체결을 위한 협상의 토대를 마련하는 것을 목적으로 한다.",
    ]),
    ("제2조 (협력 범위)", [
        "당사자들은 다음 각 호의 사항에 관하여 상호 협력한다.",
        "1. Telom-X-Gene의 안정적·지속적 공급 체계 구축",
        "2. 공급 물량, 사양(규격), 품질 기준 및 가격 조건의 협의",
        "3. 중국 내 인허가(NMPA 등) 및 등록 절차에 대한 상호 지원",
        "4. 기술자료·품질자료 등 필요 정보의 제공 및 공유",
        "5. 향후 독점/비독점 공급 여부에 관한 협의",
        "6. 기타 양 당사자가 합의하는 협력 사항",
    ]),
    ("제3조 (공급 조건의 기본 방향)", [
        "1. 대상 제품: Telom-X-Gene (상세 규격은 본계약 부속서에서 확정한다.)",
        "2. 예상 공급 물량: 연간 [ ] 단위 (협의 후 확정)",
        "3. 공급 기간: 본계약 발효일로부터 [ ]년 (갱신 가능)",
        "4. 가격 및 결제 조건: 본계약 협상 시 별도 합의하며, 통화·결제수단·인도조건(Incoterms)을 명시한다.",
        "5. 품질 보증: 갑은 합의된 품질 기준 및 관련 규제를 준수하는 제품을 공급한다.",
    ]),
    ("제4조 (인허가 및 규제 준수)", [
        "1. 을은 중국 내 본 제품의 수입·등록·판매에 필요한 인허가(국가약품감독관리국(NMPA) 등) 취득을 주도하며, "
        "갑은 이에 필요한 기술·품질 자료를 합리적 범위에서 제공한다.",
        "2. 당사자들은 각국의 수출입 규제, 의약품 관련 법규 및 국제 규범을 준수한다.",
    ]),
    ("제5조 (비밀유지)", [
        "1. 각 당사자는 본 MOU 및 협상 과정에서 알게 된 상대방의 영업비밀, 기술정보, 가격정보 등 일체의 "
        "비밀정보를 상대방의 사전 서면 동의 없이 제3자에게 공개하거나 본 MOU의 목적 외로 사용하여서는 아니 된다.",
        "2. 본 비밀유지 의무는 본 MOU 종료 후에도 [ ]년간 존속한다.",
        "3. 당사자들은 필요 시 별도의 비밀유지계약(NDA)을 체결할 수 있다.",
    ]),
    ("제6조 (법적 구속력)", [
        "1. 본 MOU는 당사자들의 협력 의사와 협상의 기본 원칙을 확인하는 문서로서, 제5조(비밀유지), "
        "제7조(독점 협상), 제9조(준거법 및 분쟁해결)를 제외한 나머지 조항은 법적 구속력을 가지지 아니한다.",
        "2. 구체적인 권리·의무 관계는 추후 체결되는 본계약에 따른다.",
    ]),
    ("제7조 (독점 협상)", [
        "당사자들은 본 MOU 체결일로부터 [ ]개월간 본 제품의 중국 내 공급과 관련하여 상호 성실히 협상하며, "
        "동 기간 중 정당한 사유 없이 제3자와 동일·유사한 거래를 위한 협상을 진행하지 아니한다.",
    ]),
    ("제8조 (유효기간)", [
        "본 MOU는 체결일로부터 효력이 발생하며, 다음 각 호 중 먼저 도래하는 시점까지 유효하다.",
        "1. 본계약 체결일",
        "2. 체결일로부터 [ ]개월이 경과한 날",
        "3. 당사자 일방이 [ ]일 전 서면 통지로 종료한 날",
    ]),
    ("제9조 (준거법 및 분쟁해결)", [
        "1. 본 MOU의 해석 및 적용에 관하여는 [대한민국 / 중국 / 제3국] 법을 준거법으로 한다.",
        "2. 본 MOU와 관련하여 발생하는 분쟁은 우선 당사자 간 협의로 해결하며, 협의로 해결되지 아니하는 경우 "
        "[싱가포르국제중재센터(SIAC) / 대한상사중재원(KCAB) / 중국국제경제무역중재위원회(CIETAC)]의 "
        "중재규칙에 따라 중재로 최종 해결한다.",
    ]),
    ("제10조 (기타)", [
        "1. 본 MOU의 수정·변경은 양 당사자의 서면 합의에 의한다.",
        "2. 본 MOU는 한국어와 [중국어/영어]로 각각 작성하며, 해석상 차이가 있는 경우 [언어]본을 우선한다.",
        "3. 본 MOU는 2부를 작성하여 각 당사자가 1부씩 보관한다.",
    ]),
]

for sec_title, bodies in sections:
    heading(sec_title, level=2)
    for b in bodies:
        para(b)

doc.add_paragraph()
para("본 MOU의 체결을 증명하기 위하여 각 당사자는 아래에 서명·날인한다.")
para("체결일: 20[ ]년 [ ]월 [ ]일", bold=True)
doc.add_paragraph()

# Signature table
sig = doc.add_table(rows=5, cols=2)
sig.style = "Table Grid"
sig.alignment = WD_TABLE_ALIGNMENT.CENTER
sig.rows[0].cells[0].text = "갑 (공급자)"
sig.rows[0].cells[1].text = "을 (구매자)"
sig.rows[1].cells[0].text = "TelomX 사"
sig.rows[1].cells[1].text = "[중국 제약회사명]"
sig.rows[2].cells[0].text = "대표자: ________________ (인)"
sig.rows[2].cells[1].text = "대표자: ________________ (인)"
sig.rows[3].cells[0].text = "직위:"
sig.rows[3].cells[1].text = "직위:"
sig.rows[4].cells[0].text = "일자:"
sig.rows[4].cells[1].text = "일자:"

# Apply east-asian font to all table cells
for tbl in (t, sig):
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Malgun Gothic"
                    run.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")

doc.add_paragraph()
note = para("※ 본 초안은 일반적 양식 예시이며, 실제 체결 전에는 반드시 양국 법률 전문가의 검토를 받으시길 권합니다.")
note.runs[0].italic = True

doc.save("MOU_TelomX_공급계약.docx")
print("saved MOU_TelomX_공급계약.docx")
