#!/usr/bin/env python3
"""Generate English and Chinese (.docx) versions of the Telom-X-Gene supply MOU."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def build_doc(cfg, font, filename):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = font
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), font)

    def set_font(run):
        run.font.name = font
        run.element.rPr.rFonts.set(qn("w:eastAsia"), font)

    def heading(text):
        h = doc.add_heading(level=2)
        run = h.add_run(text)
        set_font(run)
        run.font.color.rgb = NAVY

    def para(text, bold=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        set_font(run)
        return p

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(cfg["title"])
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = NAVY
    set_font(r)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rs = sub.add_run(cfg["subtitle"])
    rs.bold = True
    rs.font.size = Pt(13)
    set_font(rs)

    doc.add_paragraph()
    para(cfg["intro"])

    # Parties table
    t = doc.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.rows[0].cells[0].text = cfg["party_hdr"][0]
    t.rows[0].cells[1].text = cfg["party_hdr"][1]
    t.rows[1].cells[0].text = cfg["party_a"][0]
    t.rows[1].cells[1].text = cfg["party_a"][1]
    t.rows[2].cells[0].text = cfg["party_b"][0]
    t.rows[2].cells[1].text = cfg["party_b"][1]

    para(cfg["party_note"])

    for sec_title, bodies in cfg["sections"]:
        heading(sec_title)
        for b in bodies:
            para(b)

    doc.add_paragraph()
    para(cfg["closing"])
    para(cfg["date_line"], bold=True)
    doc.add_paragraph()

    sig = doc.add_table(rows=5, cols=2)
    sig.style = "Table Grid"
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (l, rr) in enumerate(cfg["sig"]):
        sig.rows[i].cells[0].text = l
        sig.rows[i].cells[1].text = rr

    for tbl in (t, sig):
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_font(run)

    doc.add_paragraph()
    note = para(cfg["disclaimer"])
    note.runs[0].italic = True

    doc.save(filename)
    print("saved", filename)


# ---------------- English ----------------
EN = {
    "title": "Memorandum of Understanding",
    "subtitle": "Memorandum of Understanding on the Supply of Telom-X-Gene",
    "intro": "This Memorandum of Understanding (the \"MOU\") is entered into between the Parties "
             "identified below for the purpose of setting out the fundamental terms concerning the "
             "supply of, and cooperation in respect of, Telom-X-Gene (the \"Product\").",
    "party_hdr": ["Category", "Party"],
    "party_a": ["Party A (Supplier)", "TelomX Inc. (Telom-X-Gene)\nAddress: [Head Office Address]\nRepresentative: [Name]"],
    "party_b": ["Party B (Purchaser)", "[Name of Chinese Pharmaceutical Company]\nAddress: [Address in China]\nRepresentative: [Name]"],
    "party_note": "(Party A and Party B are each referred to as a \"Party\" and collectively as the \"Parties\".)",
    "sections": [
        ("Preamble", [
            "Party A possesses research, development and manufacturing capabilities for Telom-X-Gene, "
            "a telomere-based gene therapy/raw-material technology, and Party B is engaged in the "
            "manufacture, distribution and sale of pharmaceuticals within the People's Republic of China "
            "(\"China\").",
            "The Parties share a mutual interest in establishing a stable supply of, and a cooperative "
            "relationship in respect of, Telom-X-Gene for the Chinese market, and enter into this MOU to "
            "confirm the basic principles and their intention to cooperate toward the conclusion of a "
            "definitive supply agreement (the \"Definitive Agreement\").",
        ]),
        ("Article 1 (Purpose)", [
            "The purpose of this MOU is to set out the basic direction and principles of cooperation "
            "necessary for Party A to supply Telom-X-Gene to Party B and for Party B to supply and utilize "
            "it in the Chinese market, and to lay the foundation for negotiating the Definitive Agreement.",
        ]),
        ("Article 2 (Scope of Cooperation)", [
            "The Parties shall cooperate on the following matters:",
            "1. Establishing a stable and continuous supply system for Telom-X-Gene;",
            "2. Negotiating supply volume, specifications, quality standards and price terms;",
            "3. Mutual support for licensing/registration procedures in China (including NMPA);",
            "4. Provision and sharing of necessary information, including technical and quality data;",
            "5. Discussion on whether the supply shall be exclusive or non-exclusive;",
            "6. Any other matters as mutually agreed by the Parties.",
        ]),
        ("Article 3 (Basic Direction of Supply Terms)", [
            "1. Product: Telom-X-Gene (detailed specifications to be finalized in an annex to the Definitive Agreement).",
            "2. Estimated Supply Volume: [ ] units per year (to be confirmed upon negotiation).",
            "3. Supply Period: [ ] years from the effective date of the Definitive Agreement (renewable).",
            "4. Price and Payment Terms: to be separately agreed during negotiation of the Definitive "
            "Agreement, specifying currency, payment method and delivery terms (Incoterms).",
            "5. Quality Assurance: Party A shall supply Products that comply with the agreed quality "
            "standards and applicable regulations.",
        ]),
        ("Article 4 (Licensing and Regulatory Compliance)", [
            "1. Party B shall lead the acquisition of the licenses/approvals required for the import, "
            "registration and sale of the Product in China (including approval by the National Medical "
            "Products Administration (NMPA)), and Party A shall provide the necessary technical and quality "
            "data to a reasonable extent.",
            "2. The Parties shall comply with applicable import/export regulations, pharmaceutical laws and "
            "international norms of each jurisdiction.",
        ]),
        ("Article 5 (Confidentiality)", [
            "1. Each Party shall not disclose to any third party, or use for any purpose other than that of "
            "this MOU, any confidential information of the other Party—including trade secrets, technical "
            "information and pricing information—learned in connection with this MOU or the negotiations, "
            "without the other Party's prior written consent.",
            "2. This confidentiality obligation shall survive for [ ] years after termination of this MOU.",
            "3. The Parties may enter into a separate non-disclosure agreement (NDA) where necessary.",
        ]),
        ("Article 6 (Legal Effect)", [
            "1. This MOU is a document confirming the Parties' intention to cooperate and the basic "
            "principles of negotiation; except for Article 5 (Confidentiality), Article 7 (Exclusive "
            "Negotiation) and Article 9 (Governing Law and Dispute Resolution), the remaining provisions "
            "shall not be legally binding.",
            "2. The specific rights and obligations shall be governed by the Definitive Agreement to be "
            "concluded subsequently.",
        ]),
        ("Article 7 (Exclusive Negotiation)", [
            "For [ ] months from the date of this MOU, the Parties shall negotiate in good faith with "
            "respect to the supply of the Product in China, and shall not, without justifiable cause, "
            "negotiate with any third party for an identical or similar transaction during such period.",
        ]),
        ("Article 8 (Term)", [
            "This MOU shall take effect on the date of execution and remain valid until the earliest of the following:",
            "1. The date of execution of the Definitive Agreement;",
            "2. The date [ ] months after the date of execution;",
            "3. The date of termination by either Party upon [ ] days' prior written notice.",
        ]),
        ("Article 9 (Governing Law and Dispute Resolution)", [
            "1. This MOU shall be governed by and construed in accordance with the laws of "
            "[the Republic of Korea / China / a third country].",
            "2. Any dispute arising in connection with this MOU shall first be resolved through mutual "
            "consultation; failing which, it shall be finally settled by arbitration under the rules of "
            "[SIAC / KCAB / CIETAC].",
        ]),
        ("Article 10 (Miscellaneous)", [
            "1. Any amendment to this MOU shall be made by written agreement of both Parties.",
            "2. This MOU is executed in Korean and [Chinese/English]; in case of any discrepancy in "
            "interpretation, the [language] version shall prevail.",
            "3. This MOU is executed in two counterparts, each Party retaining one copy.",
        ]),
    ],
    "closing": "IN WITNESS WHEREOF, the Parties have signed and sealed below to evidence the execution of this MOU.",
    "date_line": "Date of Execution: [ ] / [ ] / 20[ ]",
    "sig": [
        ("Party A (Supplier)", "Party B (Purchaser)"),
        ("TelomX Inc.", "[Name of Chinese Pharmaceutical Company]"),
        ("Representative: ________________ (Seal)", "Representative: ________________ (Seal)"),
        ("Title:", "Title:"),
        ("Date:", "Date:"),
    ],
    "disclaimer": "* This draft is a general template example. Please obtain a review by legal "
                  "professionals in both jurisdictions before execution.",
}

# ---------------- Chinese (Simplified) ----------------
ZH = {
    "title": "谅解备忘录",
    "subtitle": "关于 Telom-X-Gene 供应的谅解备忘录",
    "intro": "本谅解备忘录（以下简称\"本备忘录\"）由下列双方为约定有关 Telom-X-Gene"
             "（以下简称\"本产品\"）供应及合作的基本事项而签订。",
    "party_hdr": ["类别", "当事方"],
    "party_a": ["甲方（供应方）", "TelomX 公司（Telom-X-Gene）\n地址：[总部地址]\n法定代表人：[姓名]"],
    "party_b": ["乙方（采购方）", "[中国制药公司名称]\n地址：[中国地址]\n法定代表人：[姓名]"],
    "party_note": "（以下甲方与乙方单独称为\"一方\"，合称为\"双方\"。）",
    "sections": [
        ("前言", [
            "甲方拥有基于端粒（telomere）的基因治疗／原料技术 Telom-X-Gene 的研究、开发及生产能力；"
            "乙方在中华人民共和国（以下简称\"中国\"）境内从事药品的生产、流通及销售业务。",
            "双方对面向中国市场稳定供应 Telom-X-Gene 并建立相互合作关系具有共同意愿，"
            "为确认今后签订正式供应合同（以下简称\"正式合同\"）的基本原则与合作意向，特签订本备忘录。",
        ]),
        ("第一条（目的）", [
            "本备忘录旨在就甲方向乙方供应 Telom-X-Gene、乙方将其供应并应用于中国市场所需的合作基本方向"
            "与原则作出约定，并为今后签订正式合同的谈判奠定基础。",
        ]),
        ("第二条（合作范围）", [
            "双方就下列各项事项相互合作：",
            "1. 建立 Telom-X-Gene 稳定、持续的供应体系；",
            "2. 协商供应数量、规格、质量标准及价格条件；",
            "3. 就中国境内的许可（含 NMPA 等）及注册程序相互支持；",
            "4. 提供并共享技术资料、质量资料等必要信息；",
            "5. 就今后独家／非独家供应事宜进行协商；",
            "6. 双方约定的其他合作事项。",
        ]),
        ("第三条（供应条件的基本方向）", [
            "1. 标的产品：Telom-X-Gene（详细规格在正式合同附件中确定）。",
            "2. 预计供应数量：每年 [ ] 单位（协商后确定）。",
            "3. 供应期限：自正式合同生效之日起 [ ] 年（可续期）。",
            "4. 价格及支付条件：在正式合同谈判时另行约定，并明确币种、支付方式及交货条件（Incoterms）。",
            "5. 质量保证：甲方供应符合约定质量标准及相关法规的产品。",
        ]),
        ("第四条（许可及法规遵守）", [
            "1. 乙方主导取得本产品在中国进口、注册及销售所需的许可（含国家药品监督管理局（NMPA）等），"
            "甲方在合理范围内提供必要的技术及质量资料。",
            "2. 双方遵守各自司法管辖区的进出口法规、药品相关法律及国际规范。",
        ]),
        ("第五条（保密）", [
            "1. 各方不得未经对方事先书面同意，向第三方披露或在本备忘录目的之外使用因本备忘录或谈判过程"
            "而获悉的对方商业秘密、技术信息、价格信息等一切保密信息。",
            "2. 本保密义务在本备忘录终止后仍存续 [ ] 年。",
            "3. 双方可在必要时另行签订保密协议（NDA）。",
        ]),
        ("第六条（法律效力）", [
            "1. 本备忘录系确认双方合作意愿及谈判基本原则的文件；除第五条（保密）、第七条（独家谈判）及"
            "第九条（准据法及争议解决）外，其余条款不具有法律约束力。",
            "2. 具体权利义务关系依今后签订的正式合同确定。",
        ]),
        ("第七条（独家谈判）", [
            "自本备忘录签订之日起 [ ] 个月内，双方就本产品在中国境内的供应相互诚信谈判，"
            "且在该期间内无正当理由不得与第三方就相同或类似交易进行谈判。",
        ]),
        ("第八条（有效期）", [
            "本备忘录自签订之日起生效，至下列各项中最先到达之时点为止有效：",
            "1. 正式合同签订之日；",
            "2. 自签订之日起届满 [ ] 个月之日；",
            "3. 任何一方提前 [ ] 日书面通知终止之日。",
        ]),
        ("第九条（准据法及争议解决）", [
            "1. 本备忘录的解释与适用以 [大韩民国／中国／第三国] 法律为准据法。",
            "2. 因本备忘录产生的争议，双方应首先通过协商解决；协商不成的，"
            "提交 [新加坡国际仲裁中心（SIAC）／大韩商事仲裁院（KCAB）／中国国际经济贸易仲裁委员会（CIETAC）] "
            "按其仲裁规则仲裁解决，仲裁裁决为终局裁决。",
        ]),
        ("第十条（其他）", [
            "1. 本备忘录的修改、变更须经双方书面同意。",
            "2. 本备忘录以韩文及 [中文／英文] 分别作成；解释上存在差异时，以 [语言] 版本为准。",
            "3. 本备忘录一式两份，双方各执一份。",
        ]),
    ],
    "closing": "为证明本备忘录之签订，双方于下方签字盖章。",
    "date_line": "签订日期：20[ ] 年 [ ] 月 [ ] 日",
    "sig": [
        ("甲方（供应方）", "乙方（采购方）"),
        ("TelomX 公司", "[中国制药公司名称]"),
        ("法定代表人：________________ （盖章）", "法定代表人：________________ （盖章）"),
        ("职务：", "职务："),
        ("日期：", "日期："),
    ],
    "disclaimer": "※ 本草案为通用范本示例，签订前请务必由两国法律专业人士审阅。",
}

build_doc(EN, "Calibri", "MOU_TelomX_Supply_EN.docx")
build_doc(ZH, "SimSun", "MOU_TelomX_供应_中文.docx")
